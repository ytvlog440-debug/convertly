"""
End-to-End Verification of PDF-to-Excel and PDF-to-Word on Bank July & Aug 2026.pdf
Measures RSS memory, execution time, output file validity, and content presence.
"""

import os
import sys
import time
import asyncio
import psutil
import fitz
import openpyxl
import docx

# Ensure backend root is on sys.path
backend_dir = os.path.dirname(os.path.abspath(__file__))
if backend_dir not in sys.path:
    sys.path.insert(0, backend_dir)

from app.services.engine.pdf_to_excel import PdfToExcelConverter
from app.services.engine.office import PdfToWordConverter
from app.core.ocr_guard import get_process_rss_mb, get_active_ocr_count


async def run_test_a(pdf_path: str, output_dir: str):
    print("\n" + "=" * 70)
    print(">>> RUNNING TEST A: PDF -> Excel (Bank July & Aug 2026.pdf)")
    print("=" * 70)

    converter = PdfToExcelConverter()
    os.makedirs(output_dir, exist_ok=True)
    proc = psutil.Process()

    rss_start = proc.memory_info().rss / (1024 * 1024)
    peak_rss = rss_start
    t0 = time.perf_counter()

    # Track memory during conversion via background monitor
    done_flag = False

    async def mem_monitor():
        nonlocal peak_rss
        while not done_flag:
            current_rss = proc.memory_info().rss / (1024 * 1024)
            if current_rss > peak_rss:
                peak_rss = current_rss
            await asyncio.sleep(0.1)

    monitor_task = asyncio.create_task(mem_monitor())

    try:
        result = await converter.convert([pdf_path], output_dir, {})
    finally:
        done_flag = True
        await monitor_task

    t1 = time.perf_counter()
    rss_end = proc.memory_info().rss / (1024 * 1024)
    duration = t1 - t0

    print(f"[TEST A] Execution duration: {duration:.2f}s")
    print(f"[TEST A] RSS start: {rss_start:.1f} MB, RSS peak: {peak_rss:.1f} MB, RSS end: {rss_end:.1f} MB")
    print(f"[TEST A] Output path: {result.output_path} (size: {result.size_bytes:,} bytes)")

    assert os.path.exists(result.output_path), "Excel output file does not exist"
    assert result.size_bytes > 0, "Excel output file is empty"

    # Open and validate Excel workbook
    wb = openpyxl.load_workbook(result.output_path, data_only=True)
    sheet_names = wb.sheetnames
    print(f"[TEST A] Worksheets created: {len(sheet_names)} -> {sheet_names}")
    assert len(sheet_names) >= 8, f"Expected at least 8 sheets (non-blank pages), got {len(sheet_names)}"

    total_rows = 0
    found_transaction_keywords = set()
    keywords_to_check = {"DATE", "DETAILS", "BALANCE", "PAYMENT", "LODGEMENT", "ACCOUNT", "IRISH"}

    for name in sheet_names:
        ws = wb[name]
        total_rows += ws.max_row
        for row in ws.iter_rows(values_only=True):
            for cell in row:
                if cell and isinstance(cell, str):
                    upper_c = cell.upper()
                    for kw in keywords_to_check:
                        if kw in upper_c:
                            found_transaction_keywords.add(kw)

    print(f"[TEST A] Total rows across all sheets: {total_rows}")
    print(f"[TEST A] Found keywords in Excel: {found_transaction_keywords}")
    wb.close()

    assert total_rows > 100, f"Expected > 100 extracted rows, got {total_rows}"
    assert len(found_transaction_keywords) >= 3, f"Expected transaction keywords, found {found_transaction_keywords}"

    return {
        "duration_s": duration,
        "rss_start_mb": rss_start,
        "rss_peak_mb": peak_rss,
        "rss_end_mb": rss_end,
        "output_path": result.output_path,
        "size_bytes": result.size_bytes,
        "sheets": len(sheet_names),
        "total_rows": total_rows,
        "keywords": list(found_transaction_keywords),
    }


async def run_test_b(pdf_path: str, output_dir: str):
    print("\n" + "=" * 70)
    print(">>> RUNNING TEST B: PDF -> Word (Bank July & Aug 2026.pdf)")
    print("=" * 70)

    converter = PdfToWordConverter()
    os.makedirs(output_dir, exist_ok=True)
    proc = psutil.Process()

    rss_start = proc.memory_info().rss / (1024 * 1024)
    peak_rss = rss_start
    t0 = time.perf_counter()

    done_flag = False

    async def mem_monitor():
        nonlocal peak_rss
        while not done_flag:
            current_rss = proc.memory_info().rss / (1024 * 1024)
            if current_rss > peak_rss:
                peak_rss = current_rss
            await asyncio.sleep(0.1)

    monitor_task = asyncio.create_task(mem_monitor())

    try:
        result = await converter.convert([pdf_path], output_dir, {})
    finally:
        done_flag = True
        await monitor_task

    t1 = time.perf_counter()
    rss_end = proc.memory_info().rss / (1024 * 1024)
    duration = t1 - t0

    print(f"[TEST B] Execution duration: {duration:.2f}s")
    print(f"[TEST B] RSS start: {rss_start:.1f} MB, RSS peak: {peak_rss:.1f} MB, RSS end: {rss_end:.1f} MB")
    print(f"[TEST B] Output path: {result.output_path} (size: {result.size_bytes:,} bytes)")

    assert os.path.exists(result.output_path), "Word output file does not exist"
    assert result.size_bytes > 0, "Word output file is empty"

    # Open and validate Word docx
    doc = docx.Document(result.output_path)
    paragraphs = doc.paragraphs
    print(f"[TEST B] Paragraphs created: {len(paragraphs)}")
    assert len(paragraphs) > 10, f"Expected > 10 paragraphs, got {len(paragraphs)}"

    full_text = "\n".join(p.text for p in paragraphs)
    print(f"[TEST B] Extracted text character count: {len(full_text):,}")

    found_transaction_keywords = set()
    keywords_to_check = {"BANK", "DATE", "BALANCE", "PAYMENT", "ACCOUNT", "STATEMENT"}
    upper_full = full_text.upper()
    for kw in keywords_to_check:
        if kw in upper_full:
            found_transaction_keywords.add(kw)

    print(f"[TEST B] Found keywords in Word: {found_transaction_keywords}")
    assert len(found_transaction_keywords) >= 3, f"Expected transaction keywords, found {found_transaction_keywords}"

    return {
        "duration_s": duration,
        "rss_start_mb": rss_start,
        "rss_peak_mb": peak_rss,
        "rss_end_mb": rss_end,
        "output_path": result.output_path,
        "size_bytes": result.size_bytes,
        "paragraphs": len(paragraphs),
        "char_count": len(full_text),
        "keywords": list(found_transaction_keywords),
    }


async def main():
    pdf_path = os.path.abspath(r"D:\convertlytools\Bank July & Aug 2026.pdf")
    output_dir = os.path.abspath(r"D:\convertlytools\backend\temp_e2e_outputs")

    print(f"Target PDF: {pdf_path}")
    assert os.path.exists(pdf_path), f"File not found: {pdf_path}"

    res_a = await run_test_a(pdf_path, output_dir)
    res_b = await run_test_b(pdf_path, output_dir)

    print("\n" + "=" * 70)
    print("ALL TESTS COMPLETED SUCCESSFULLY!")
    print(f"TEST A (Excel): Peak RSS = {res_a['rss_peak_mb']:.1f} MB, Duration = {res_a['duration_s']:.2f}s, Sheets = {res_a['sheets']}, Rows = {res_a['total_rows']}")
    print(f"TEST B (Word):  Peak RSS = {res_b['rss_peak_mb']:.1f} MB, Duration = {res_b['duration_s']:.2f}s, Paragraphs = {res_b['paragraphs']}, Chars = {res_b['char_count']}")
    print("=" * 70)


if __name__ == "__main__":
    asyncio.run(main())
