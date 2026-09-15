"""
Convertly V2 — Enterprise OpenXML Excel Writer
Synthesizes structured TableBlocks, KeyValueBlocks, and cell spans into production XLSX files.
Features:
  1. Native ws.merge_cells() for multi-column and multi-row spans (timetables, balance sheets)
  2. Semantic cell typing and localized number/currency formats
  3. Hierarchical cell indentation for financial statements
  4. Top-level metadata and footer totals placement
  5. Multi-page consolidation into 'All Data (Consolidated)' master sheet
  6. Auto-fitting column widths with comfortable padding
"""

import os
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from typing import List, Dict, Any, Optional
from datetime import datetime

from app.services.engine.table_extractor.models import PageLayout, TableBlock, KeyValueBlock, TableCell


class ExcelDocumentWriter:
    """
    Renders extracted PageLayout objects into an enterprise-grade XLSX workbook.
    """

    def __init__(self):
        self.header_font = Font(name='Calibri', size=11, bold=True, color='0F172A')
        self.header_fill = PatternFill(start_color='F1F5F9', end_color='F1F5F9', fill_type='solid')
        self.header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

        self.data_font = Font(name='Calibri', size=11, color='1E293B')
        self.num_alignment = Alignment(horizontal='right', vertical='center')
        self.date_alignment = Alignment(horizontal='center', vertical='center')
        self.text_alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)

        self.thin_border = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )

        self.meta_key_font = Font(name='Calibri', size=10, bold=True, color='475569')
        self.meta_val_font = Font(name='Calibri', size=10, color='0F172A')

    def write_workbook(self, page_layouts: List[PageLayout], output_path: str) -> Dict[str, Any]:
        """
        Creates an XLSX file from extracted page layouts.
        """
        wb = openpyxl.Workbook()
        default_sheet = wb.active
        if default_sheet:
            wb.remove(default_sheet)

        total_tables = sum(len(p.tables) for p in page_layouts)
        total_rows_written = sum(sum(len(t.rows) for t in p.tables) for p in page_layouts)
        pages_with_data = 0

        # Check for multi-page consolidation eligibility
        valid_pages = [p for p in page_layouts if p.tables]
        can_consolidate = len(valid_pages) >= 2
        first_cols = None

        if can_consolidate:
            for p in valid_pages:
                primary = p.tables[0]
                if not primary.rows:
                    can_consolidate = False
                    break
                cols = primary.col_count
                if first_cols is None:
                    first_cols = cols
                elif first_cols != cols:
                    can_consolidate = False
                    break

        # 1. Build Consolidated Master Sheet if eligible
        if can_consolidate and first_cols and first_cols >= 2:
            ws_cons = wb.create_sheet(title="All Data (Consolidated)")
            cons_curr_row = 1
            for idx, p in enumerate(valid_pages):
                primary = p.tables[0]
                rows_to_render = primary.rows if idx == 0 else (primary.rows[1:] if primary.rows[0].is_header else primary.rows)
                cons_curr_row = self._render_table_to_sheet(ws_cons, rows_to_render, cons_curr_row)
            self._auto_fit_columns(ws_cons)

        # 2. Build Individual Page Worksheets
        for p in page_layouts:
            if not p.tables and not p.key_values:
                continue

            pages_with_data += 1
            ws_title = f"Page {p.page_idx + 1}"
            ws = wb.create_sheet(title=ws_title[:31])
            curr_row = 1

            # Render top key-value block (e.g. Invoice #, Bill-To, Date)
            header_kvs = [kv for kv in p.key_values if kv.section == "header"]
            for kv_block in header_kvs:
                curr_row = self._render_key_value_block(ws, kv_block, curr_row)
                curr_row += 1  # 1-row gap

            # Render tables on page
            for t_idx, table in enumerate(p.tables):
                if t_idx > 0:
                    curr_row += 2  # 2-row gap between tables
                curr_row = self._render_table_to_sheet(ws, table.rows, curr_row)

            # Render bottom footer totals / summary block
            footer_kvs = [kv for kv in p.key_values if kv.section == "footer"]
            for kv_block in footer_kvs:
                curr_row += 1  # 1-row gap
                curr_row = self._render_key_value_block(ws, kv_block, curr_row, align_right=True)

            self._auto_fit_columns(ws)

        # Fallback if no tables/data detected
        if not wb.sheetnames:
            ws = wb.create_sheet(title="Results")
            ws.cell(row=1, column=1, value="No structured tables found in the document.")
            ws.cell(row=2, column=1, value="The PDF may contain only graphics, forms without lines, or unformatted text.")
            ws.column_dimensions['A'].width = 65

        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
        wb.save(output_path)


        return {
            "total_tables_extracted": total_tables,
            "total_rows": total_rows_written,
            "pages_with_data": pages_with_data,
            "total_pages_processed": len(page_layouts),
            "worksheets_created": len(wb.sheetnames),
        }

    def _render_table_to_sheet(
        self,
        ws: openpyxl.worksheet.worksheet.Worksheet,
        rows: List[Any],
        start_row: int
    ) -> int:
        """Writes rows and handles colspans/rowspans via ws.merge_cells."""
        curr_row = start_row
        merges_to_apply = []
        merged_covered = set()

        # Step 1: Discover all merged ranges and mark secondary cells as covered
        for r_offset, row in enumerate(rows):
            target_r = curr_row + r_offset
            for cell in row.cells:
                target_c = cell.col_idx + 1
                if cell.colspan > 1 or cell.rowspan > 1:
                    end_r = target_r + cell.rowspan - 1
                    end_c = target_c + cell.colspan - 1
                    is_hdr = cell.is_header or row.is_header
                    merges_to_apply.append((target_r, target_c, end_r, end_c, is_hdr))
                    for mr in range(target_r, end_r + 1):
                        for mc in range(target_c, end_c + 1):
                            if (mr, mc) != (target_r, target_c):
                                merged_covered.add((mr, mc))

        # Step 2: Populate values, fonts, fills, and alignments
        for r_offset, row in enumerate(rows):
            target_r = curr_row + r_offset
            for cell in row.cells:
                target_c = cell.col_idx + 1

                c = ws.cell(row=target_r, column=target_c)
                c.border = self.thin_border

                if cell.is_header or row.is_header:
                    c.font = self.header_font
                    c.fill = self.header_fill
                    c.alignment = self.header_alignment
                else:
                    c.font = self.data_font
                    # Alignment heuristics
                    if cell.data_type in ("currency", "percentage", "number"):
                        c.alignment = self.num_alignment
                    elif cell.data_type == "date":
                        c.alignment = self.date_alignment
                    else:
                        if cell.indent_level > 0:
                            c.alignment = Alignment(horizontal='left', vertical='center', indent=cell.indent_level, wrap_text=True)
                        else:
                            c.alignment = self.text_alignment

                if cell.format_code:
                    c.number_format = cell.format_code

                # Only write values into top-left origin cells
                if (target_r, target_c) not in merged_covered:
                    c.value = cell.typed_value if cell.typed_value is not None else cell.text
                else:
                    c.value = None

        # Step 3: Format all cells within each merged range for complete borders & fill
        for start_r, start_c, end_r, end_c, is_hdr in merges_to_apply:
            for mr in range(start_r, end_r + 1):
                for mc in range(start_c, end_c + 1):
                    cell_in_box = ws.cell(row=mr, column=mc)
                    cell_in_box.border = self.thin_border
                    if is_hdr:
                        cell_in_box.fill = self.header_fill

        # Step 4: Apply openpyxl ws.merge_cells
        for start_r, start_c, end_r, end_c, is_hdr in merges_to_apply:
            ws.merge_cells(
                start_row=start_r,
                start_column=start_c,
                end_row=end_r,
                end_column=end_c
            )
            # Ensure text in spanning header is centered
            top_left = ws.cell(row=start_r, column=start_c)
            if is_hdr or end_c > start_c:
                top_left.alignment = self.header_alignment

        return curr_row + len(rows)


    def _render_key_value_block(
        self,
        ws: openpyxl.worksheet.worksheet.Worksheet,
        kv_block: KeyValueBlock,
        start_row: int,
        align_right: bool = False
    ) -> int:
        """Renders 2-column key-value pairs cleanly."""
        curr_row = start_row
        col_offset = max(1, (ws.max_column or 2) - 2) if align_right else 1

        for key, val in kv_block.items:
            cell_k = ws.cell(row=curr_row, column=col_offset, value=f"{key}:")
            cell_k.font = self.meta_key_font
            cell_k.alignment = Alignment(horizontal='right' if align_right else 'left', vertical='center')

            cell_v = ws.cell(row=curr_row, column=col_offset + 1, value=val)
            cell_v.font = self.meta_val_font
            cell_v.alignment = Alignment(horizontal='right' if align_right else 'left', vertical='center')

            curr_row += 1

        return curr_row

    def _auto_fit_columns(self, ws: openpyxl.worksheet.worksheet.Worksheet):
        """Calculates optimal column width with padding."""
        for col_idx in range(1, (ws.max_column or 0) + 1):
            col_letter = get_column_letter(col_idx)
            max_len = 10
            for row_idx in range(1, (ws.max_row or 0) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                if cell.value is not None:
                    lines = str(cell.value).split('\n')
                    longest = max(len(l) for l in lines) if lines else 0
                    max_len = max(max_len, min(longest + 3, 60))
            ws.column_dimensions[col_letter].width = max_len
