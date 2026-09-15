"""
Convertly V2 — PDF to Excel Converter Engine
Extracts tables, financial data, and structured content from PDF documents
into editable Excel (XLSX) spreadsheets with smart type detection.

Pipeline:
  1. PyMuPDF validates and opens the PDF
  2. Scanned PDF auto-detection via text density heuristic
  3. OCR fallback via pytesseract for scanned pages
  4. pdfplumber (primary) + tabula-py (secondary) table extraction
  5. Smart data type detection (currency, dates, percentages, numbers)
  6. openpyxl XLSX generation with formatting and auto-width columns
"""

import asyncio
import io
import os
import re
import uuid
from datetime import datetime
from typing import List, Dict, Any, Optional, Tuple

import fitz  # PyMuPDF
import openpyxl
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from app.core.errors import FileValidationError, ConversionExecutionError
from app.core.logging import logger
from app.services.engine.base import BaseConverter, ConversionResult
from app.services.engine.pdf import open_and_validate_pdf
from app.services.engine.table_extractor import EnterpriseTableExtractor


# ---------------------------------------------------------------------------
# Data Type Detection Utilities
# ---------------------------------------------------------------------------

# Currency symbols and patterns
_CURRENCY_SYMBOLS = r'[\$\u20ac\u00a3\u00a5\u20b9\u20a9\u20bd\u20ba\u20b1\u20bf]'
_CURRENCY_PATTERN = re.compile(
    rf'^\s*({_CURRENCY_SYMBOLS}|CHF|USD|EUR|GBP|JPY|CAD|AUD)\s*[\-\(]?\s*[\d,.]+\s*\)?\s*$'
    r'|'
    rf'^\s*[\-\(]?\s*[\d,.]+\s*\)?\s*({_CURRENCY_SYMBOLS}|CHF|USD|EUR|GBP|JPY|CAD|AUD)\s*$',
    re.UNICODE | re.IGNORECASE
)

_CURRENCY_FORMAT_MAP = {
    '$': '$#,##0.00;($#,##0.00);"-"',
    '€': '€#,##0.00;(€#,##0.00);"-"',
    '£': '£#,##0.00;(£#,##0.00);"-"',
    '¥': '¥#,##0;(¥#,##0);"-"',
    '₹': '₹#,##0.00;(₹#,##0.00);"-"',
    '₩': '₩#,##0;(₩#,##0);"-"',
    '₽': '₽#,##0.00;(₽#,##0.00);"-"',
    '₺': '₺#,##0.00;(₺#,##0.00);"-"',
    'USD': '$#,##0.00;($#,##0.00);"-"',
    'EUR': '€#,##0.00;(€#,##0.00);"-"',
    'GBP': '£#,##0.00;(£#,##0.00);"-"',
    'JPY': '¥#,##0;(¥#,##0);"-"',
    'CHF': 'CHF #,##0.00;(CHF #,##0.00);"-"',
    'CAD': '$#,##0.00;($#,##0.00);"-"',
    'AUD': '$#,##0.00;($#,##0.00);"-"',
}

# Accounting negative: (1,234.56)
_ACCOUNTING_NEG_PATTERN = re.compile(
    r'^\s*\(\s*[\d,.]+\s*\)\s*$'
)

# Percentage: 12.5%, 0.5 %, -3.2%
_PERCENTAGE_PATTERN = re.compile(
    r'^\s*[\-+]?\s*[\d,.]+\s*%\s*$'
)

# Date patterns
_DATE_PATTERNS = [
    # MM/DD/YYYY, MM-DD-YYYY, MM.DD.YYYY
    (re.compile(r'^\s*(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{4})\s*$'), '%m/%d/%Y'),
    # DD/MM/YYYY (European)
    (re.compile(r'^\s*(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{4})\s*$'), '%d/%m/%Y'),
    # YYYY-MM-DD (ISO)
    (re.compile(r'^\s*(\d{4})[/\-.](\d{1,2})[/\-.](\d{1,2})\s*$'), '%Y-%m-%d'),
    # Mon DD, YYYY or DD Mon YYYY
    (re.compile(r'^\s*\w{3,9}\s+\d{1,2},?\s+\d{4}\s*$'), None),
    (re.compile(r'^\s*\d{1,2}\s+\w{3,9}\s+\d{4}\s*$'), None),
]

# Pure number: 1,234.56  or  -1234.56  or  1234  or  1.234,56 (European)
_NUMBER_PATTERN = re.compile(
    r'^\s*[\-+]?\s*[\d,.]+\s*$'
)


def _detect_currency_symbol(text: str) -> Tuple[str, str]:
    """Extract currency symbol/code and cleaned numerical text."""
    for symbol, fmt in _CURRENCY_FORMAT_MAP.items():
        if symbol in text or symbol.lower() in text.lower():
            cleaned = re.sub(re.escape(symbol), '', text, flags=re.IGNORECASE).strip()
            return symbol, cleaned
    # Fallback to default USD symbol
    cleaned = re.sub(rf'{_CURRENCY_SYMBOLS}', '', text).strip()
    return '$', cleaned


def _parse_number(text: str) -> Optional[float]:
    """
    Parse a string as a number, handling accounting parentheses (1,234.56),
    US standard comma-thousands (1,234.56), and European dot-thousands (1.234,56).
    """
    cleaned = text.strip()
    if not cleaned:
        return None

    # Handle accounting negative: (1,234.56) -> -1234.56
    is_negative = False
    if _ACCOUNTING_NEG_PATTERN.match(cleaned):
        is_negative = True
        cleaned = cleaned.replace('(', '').replace(')', '').strip()

    # Handle explicit negative sign
    if cleaned.startswith('-') or cleaned.startswith('\u2212'):
        is_negative = True
        cleaned = cleaned.lstrip('-\u2212').strip()
    elif cleaned.startswith('+'):
        cleaned = cleaned.lstrip('+').strip()

    # Detect European formatting: e.g. 1.234,56 (dot thousands, comma decimal)
    # or simple comma decimals with 1-2 digits where no dot exists and not 3-digit thousands (e.g. 125,50)
    has_dot_thousands = bool(re.search(r'^\d{1,3}(\.\d{3})+,\d+$', cleaned))
    has_comma_decimals = bool(
        ',' in cleaned and '.' not in cleaned and
        re.search(r',\d{1,2}$', cleaned) and
        not re.search(r'^\d{1,3}(,\d{3})+$', cleaned)
    )

    if has_dot_thousands or has_comma_decimals:
        cleaned = cleaned.replace('.', '').replace(',', '.')
    else:
        # Standard US/UK format: remove thousands comma
        cleaned = cleaned.replace(',', '')

    if not cleaned:
        return None

    try:
        value = float(cleaned)
        return -value if is_negative else value
    except (ValueError, OverflowError):
        return None


def detect_cell_type(text: str) -> Tuple[Any, Optional[str]]:
    """
    Detect the data type of a cell value and return (typed_value, number_format).
    Returns (original_text, None) if no specific type is detected.
    Automatically filters XML control characters that crash openpyxl.
    """
    if text is None:
        return None, None

    # Clean illegal XML characters
    raw_str = ILLEGAL_CHARACTERS_RE.sub('', str(text)).strip()
    if not raw_str:
        return None, None

    # 1. Currency detection (e.g. $1,250.00, €450.50, £99.00, (¥10,000))
    if _CURRENCY_PATTERN.match(raw_str):
        symbol, num_str = _detect_currency_symbol(raw_str)
        value = _parse_number(num_str)
        if value is not None:
            fmt = _CURRENCY_FORMAT_MAP.get(symbol, f'{symbol}#,##0.00;({symbol}#,##0.00);"-"')
            return value, fmt

    # 2. Percentage detection (e.g. 12.5%, -3.2%)
    if _PERCENTAGE_PATTERN.match(raw_str):
        num_str = raw_str.replace('%', '').strip()
        value = _parse_number(num_str)
        if value is not None:
            return value / 100.0, '0.00%'

    # 3. Date detection (ISO, US, European, Textual)
    for pattern, _ in _DATE_PATTERNS:
        if pattern.match(raw_str):
            for try_fmt in [
                '%Y-%m-%d', '%m/%d/%Y', '%d/%m/%Y', '%Y/%m/%d',
                '%m-%d-%Y', '%d-%m-%Y', '%m.%d.%Y', '%d.%m.%Y',
                '%b %d, %Y', '%B %d, %Y', '%d %b %Y', '%d %B %Y',
                '%b %d %Y', '%B %d %Y',
            ]:
                try:
                    dt = datetime.strptime(raw_str, try_fmt)
                    if 1900 <= dt.year <= 2100:
                        return dt, 'YYYY-MM-DD'
                except ValueError:
                    continue
            break

    # 4. Accounting negative: (1,234.56)
    if _ACCOUNTING_NEG_PATTERN.match(raw_str):
        value = _parse_number(raw_str)
        if value is not None:
            return value, '#,##0.00;(#,##0.00);"-"'

    # 5. Plain number (integer or floating point)
    if _NUMBER_PATTERN.match(raw_str):
        value = _parse_number(raw_str)
        if value is not None:
            # Whole numbers without decimals
            if value == int(value) and abs(value) < 1e15 and '.' not in raw_str and ',' not in raw_str[-3:]:
                return int(value), '#,##0'
            return value, '#,##0.00'

    # 6. Default: return sanitized string
    return raw_str, None


# ---------------------------------------------------------------------------
# Table Extraction Strategies
# ---------------------------------------------------------------------------

def _extract_tables_pymupdf(pdf_path: str, page_indices: List[int]) -> Dict[int, List[List[List[str]]]]:
    """
    Fallback extraction using PyMuPDF's built-in table finder and layout heuristics.
    Requires no external dependencies beyond fitz/pymupdf.
    """
    results: Dict[int, List[List[List[str]]]] = {}
    try:
        doc = fitz.open(pdf_path)
        for page_idx in page_indices:
            if page_idx >= len(doc):
                continue
            page = doc[page_idx]
            page_tables: List[List[List[str]]] = []

            # Try PyMuPDF find_tables (available in PyMuPDF >= 1.23.0)
            try:
                tabs = page.find_tables()
                if tabs and tabs.tables:
                    for tab in tabs.tables:
                        extracted = tab.extract()
                        if extracted:
                            cleaned_table = []
                            for row in extracted:
                                cleaned_row = [
                                    (cell.strip() if cell and isinstance(cell, str) else (str(cell).strip() if cell is not None else ''))
                                    for cell in row
                                ]
                                cleaned_table.append(cleaned_row)
                            if cleaned_table:
                                page_tables.append(cleaned_table)
            except Exception as e:
                logger.debug(f"PyMuPDF find_tables warning on page {page_idx}: {e}")

            # Fallback to text line grid heuristics
            if not page_tables:
                text = page.get_text("text")
                if text and text.strip():
                    lines = [line.strip() for line in text.strip().split('\n') if line.strip()]
                    grid = []
                    for line in lines:
                        cells = re.split(r'\t|  {2,}', line)
                        cells = [c.strip() for c in cells if c.strip()]
                        if cells:
                            grid.append(cells)
                    if grid:
                        page_tables.append(grid)

            if page_tables:
                results[page_idx] = page_tables
        doc.close()
    except Exception as e:
        logger.warning(f"PyMuPDF table extraction failed: {e}")

    return results


def _extract_tables_pdfplumber(pdf_path: str, page_indices: List[int]) -> Dict[int, List[List[List[str]]]]:
    """
    Primary extraction using pdfplumber with automatic fallback to PyMuPDF.
    Uses multi-strategy table detection:
      1. Bordered tables (lines & lines)
      2. Semi-bordered tables with horizontal dividers (common in financial statements & invoices)
      3. Borderless tables (text & text)
      4. Text & surrounding key-value metadata preservation (Invoice headers & summary totals)
    Returns: {page_index: [table1_rows, table2_rows, ...]}
    """
    try:
        import pdfplumber
    except ImportError:
        logger.info("pdfplumber not installed, using native PyMuPDF table extraction")
        return _extract_tables_pymupdf(pdf_path, page_indices)

    results: Dict[int, List[List[List[str]]]] = {}

    try:
        with pdfplumber.open(pdf_path) as pdf:
            for page_idx in page_indices:
                if page_idx >= len(pdf.pages):
                    continue

                page = pdf.pages[page_idx]
                page_tables: List[List[List[str]]] = []

                # Find tables using bounding-box aware finder
                # Strategy 1: Explicit line-based detection (bordered tables)
                found_tables = page.find_tables(table_settings={
                    "vertical_strategy": "lines",
                    "horizontal_strategy": "lines",
                    "snap_tolerance": 4,
                    "join_tolerance": 4,
                    "edge_min_length": 10,
                })

                # Strategy 2: Text verticals with line horizontals (very common in invoices & bank statements)
                if not found_tables:
                    found_tables = page.find_tables(table_settings={
                        "vertical_strategy": "text",
                        "horizontal_strategy": "lines",
                        "snap_tolerance": 5,
                        "join_tolerance": 5,
                        "min_words_vertical": 2,
                        "min_words_horizontal": 1,
                    })

                # Strategy 3: Text-only borderless tables
                if not found_tables:
                    found_tables = page.find_tables(table_settings={
                        "vertical_strategy": "text",
                        "horizontal_strategy": "text",
                        "snap_tolerance": 5,
                        "join_tolerance": 5,
                        "min_words_vertical": 2,
                        "min_words_horizontal": 1,
                    })

                if found_tables:
                    # Sort tables top-to-bottom by y-coordinate
                    sorted_tables = sorted(found_tables, key=lambda t: t.bbox[1])

                    # Check for header metadata above the first table (e.g. Invoice #, Bill-To, Date)
                    first_top = sorted_tables[0].bbox[1]
                    if first_top > 45:
                        try:
                            above_crop = page.crop((0, 0, page.width, max(0, first_top - 4)))
                            above_text = above_crop.extract_text()
                            if above_text and above_text.strip():
                                meta_rows = []
                                for line in above_text.strip().split('\n'):
                                    cleaned_line = line.strip()
                                    if cleaned_line:
                                        parts = re.split(r'\t|  {2,}|:\s+', cleaned_line, maxsplit=2)
                                        parts = [p.strip() for p in parts if p.strip()]
                                        if parts:
                                            meta_rows.append(parts)
                                if meta_rows:
                                    page_tables.append(meta_rows)
                        except Exception as e:
                            logger.debug(f"Above-table metadata extraction warning: {e}")

                    # Extract the structured table bodies
                    for t in sorted_tables:
                        raw_extracted = t.extract()
                        if raw_extracted and len(raw_extracted) > 0:
                            cleaned_table = []
                            for row in raw_extracted:
                                cleaned_row = [
                                    (cell.strip() if cell and isinstance(cell, str) else (str(cell).strip() if cell is not None else ''))
                                    for cell in row
                                ]
                                cleaned_table.append(cleaned_row)
                            if cleaned_table:
                                page_tables.append(cleaned_table)

                    # Check for summary metadata below the last table (e.g. Subtotals, Taxes, Total Due, Terms)
                    last_bottom = sorted_tables[-1].bbox[3]
                    if last_bottom < page.height - 35:
                        try:
                            below_crop = page.crop((0, min(page.height, last_bottom + 4), page.width, page.height))
                            below_text = below_crop.extract_text()
                            if below_text and below_text.strip():
                                summary_rows = []
                                for line in below_text.strip().split('\n'):
                                    cleaned_line = line.strip()
                                    if cleaned_line:
                                        parts = re.split(r'\t|  {2,}|:\s+', cleaned_line, maxsplit=2)
                                        parts = [p.strip() for p in parts if p.strip()]
                                        if parts:
                                            summary_rows.append(parts)
                                if summary_rows:
                                    page_tables.append(summary_rows)
                        except Exception as e:
                            logger.debug(f"Below-table metadata extraction warning: {e}")

                # Strategy 4: Fallback to full-page text grid if no tables detected
                if not page_tables:
                    text = page.extract_text()
                    if text and text.strip():
                        lines = [line.strip() for line in text.strip().split('\n') if line.strip()]
                        if lines:
                            grid = []
                            for line in lines:
                                cells = re.split(r'\t|  {2,}', line)
                                cells = [c.strip() for c in cells if c.strip()]
                                if cells:
                                    grid.append(cells)
                            if grid:
                                page_tables.append(grid)

                if page_tables:
                    results[page_idx] = page_tables

    except Exception as e:
        logger.warning(f"pdfplumber extraction failed: {e}")

    return results


def _extract_tables_tabula(pdf_path: str, page_indices: List[int]) -> Dict[int, List[List[List[str]]]]:
    """
    Secondary extraction using tabula-py (requires Java).
    Returns same format as pdfplumber extractor.
    """
    results: Dict[int, List[List[List[str]]]] = {}

    try:
        import tabula

        for page_idx in page_indices:
            page_num = page_idx + 1  # tabula uses 1-based page numbers
            page_tables: List[List[List[str]]] = []

            # Try lattice mode first (for bordered tables)
            try:
                dfs = tabula.read_pdf(
                    pdf_path,
                    pages=str(page_num),
                    lattice=True,
                    multiple_tables=True,
                    pandas_options={'header': None},
                    silent=True,
                )
                for df in dfs:
                    if df is not None and not df.empty:
                        table = []
                        for _, row in df.iterrows():
                            table.append([str(cell).strip() if str(cell).strip() != 'nan' else '' for cell in row])
                        if table:
                            page_tables.append(table)
            except Exception:
                pass

            # If lattice didn't find tables, try stream mode (borderless)
            if not page_tables:
                try:
                    dfs = tabula.read_pdf(
                        pdf_path,
                        pages=str(page_num),
                        stream=True,
                        multiple_tables=True,
                        pandas_options={'header': None},
                        silent=True,
                    )
                    for df in dfs:
                        if df is not None and not df.empty:
                            table = []
                            for _, row in df.iterrows():
                                table.append([str(cell).strip() if str(cell).strip() != 'nan' else '' for cell in row])
                            if table:
                                page_tables.append(table)
                except Exception:
                    pass

            if page_tables:
                results[page_idx] = page_tables

    except ImportError:
        logger.info("tabula-py not available, skipping tabula extraction")
    except Exception as e:
        logger.warning(f"tabula extraction failed: {e}")

    return results


def _extract_tables_ocr(
    doc: fitz.Document,
    page_indices: List[int],
    tesseract_bin: Optional[str] = None,
) -> Dict[int, List[List[List[str]]]]:
    """
    OCR-based table extraction for scanned PDFs.
    Renders each page, runs OCR, then attempts grid reconstruction
    from the OCR text output.
    """
    results: Dict[int, List[List[List[str]]]] = {}

    if not tesseract_bin:
        return results

    try:
        import pytesseract
        from PIL import Image

        pytesseract.pytesseract.tesseract_cmd = tesseract_bin
    except Exception as e:
        logger.warning(f"Failed to initialize pytesseract for table OCR: {e}")
        return results

    for page_idx in page_indices:
        if page_idx >= len(doc):
            continue

        page = doc[page_idx]
        page_tables: List[List[List[str]]] = []

        try:
            # Render at 200 DPI for good OCR quality
            pix = page.get_pixmap(dpi=200)
            pil_img = Image.open(io.BytesIO(pix.tobytes("png")))

            # Use pytesseract TSV output for coordinate-based cell detection
            try:
                tsv_data = pytesseract.image_to_data(
                    pil_img, lang="eng", output_type=pytesseract.Output.DICT, timeout=15
                )

                # Group words by their block/paragraph/line
                word_entries = []
                for i in range(len(tsv_data['text'])):
                    text = str(tsv_data['text'][i]).strip()
                    if text and tsv_data['conf'][i] > 20:  # Confidence threshold
                        word_entries.append({
                            'text': text,
                            'left': tsv_data['left'][i],
                            'top': tsv_data['top'][i],
                            'width': tsv_data['width'][i],
                            'height': tsv_data['height'][i],
                            'block_num': tsv_data['block_num'][i],
                            'line_num': tsv_data['line_num'][i],
                        })

                if word_entries:
                    # Group by y-coordinate proximity (same row)
                    rows_by_y: Dict[int, List[Dict]] = {}
                    y_tolerance = 12  # pixels

                    for entry in word_entries:
                        center_y = entry['top'] + entry['height'] // 2
                        matched_y = None
                        for existing_y in rows_by_y:
                            if abs(center_y - existing_y) <= y_tolerance:
                                matched_y = existing_y
                                break
                        if matched_y is not None:
                            rows_by_y[matched_y].append(entry)
                        else:
                            rows_by_y[center_y] = [entry]

                    # Sort rows by y-coordinate
                    sorted_ys = sorted(rows_by_y.keys())

                    # Detect columns by analyzing x-coordinate clusters
                    all_x_positions = []
                    for y_key in sorted_ys:
                        for entry in rows_by_y[y_key]:
                            all_x_positions.append(entry['left'])

                    if all_x_positions:
                        # Find column boundaries via gap detection
                        unique_xs = sorted(set(all_x_positions))
                        col_boundaries = [0]

                        if len(unique_xs) > 1:
                            gaps = []
                            for i in range(1, len(unique_xs)):
                                gaps.append((unique_xs[i] - unique_xs[i - 1], unique_xs[i]))

                            # Significant gaps define column boundaries
                            avg_gap = sum(g[0] for g in gaps) / len(gaps) if gaps else 50
                            threshold = max(avg_gap * 1.5, 30)

                            for gap_size, gap_x in gaps:
                                if gap_size > threshold:
                                    col_boundaries.append(gap_x)

                        col_boundaries.append(99999)  # Right edge

                        # Build grid
                        grid = []
                        for y_key in sorted_ys:
                            row_words = sorted(rows_by_y[y_key], key=lambda e: e['left'])
                            row_cells = [''] * (len(col_boundaries) - 1)

                            for entry in row_words:
                                # Find which column this word belongs to
                                for col_idx in range(len(col_boundaries) - 1):
                                    if col_boundaries[col_idx] <= entry['left'] < col_boundaries[col_idx + 1]:
                                        if row_cells[col_idx]:
                                            row_cells[col_idx] += ' ' + entry['text']
                                        else:
                                            row_cells[col_idx] = entry['text']
                                        break

                            grid.append(row_cells)

                        if grid:
                            page_tables.append(grid)

            except Exception as e:
                logger.warning(f"OCR TSV extraction failed for page {page_idx + 1}: {e}")
                # Fallback: simple line-based extraction
                text = pytesseract.image_to_string(pil_img, lang="eng", timeout=15)
                if text and text.strip():
                    lines = [line.strip() for line in text.strip().split('\n') if line.strip()]
                    grid = []
                    for line in lines:
                        cells = re.split(r'\t|  {2,}', line)
                        cells = [c.strip() for c in cells if c.strip()]
                        if cells:
                            grid.append(cells)
                    if grid:
                        page_tables.append(grid)

        except Exception as e:
            logger.warning(f"OCR page rendering failed for page {page_idx + 1}: {e}")

        if page_tables:
            results[page_idx] = page_tables

    return results


# ---------------------------------------------------------------------------
# XLSX Generation
# ---------------------------------------------------------------------------

def _normalize_table_columns(table: List[List[str]]) -> List[List[str]]:
    """Ensure all rows in a table have the same number of columns."""
    if not table:
        return table
    max_cols = max(len(row) for row in table)
    return [row + [''] * (max_cols - len(row)) for row in table]


def _is_likely_header(row: List[str], total_rows: int) -> bool:
    """Heuristic: a header row usually has mostly non-numeric text cells."""
    if total_rows < 2:
        return False
    non_empty = [c for c in row if c and str(c).strip()]
    if not non_empty:
        return False
    numeric_count = sum(1 for c in non_empty if _NUMBER_PATTERN.match(str(c).strip()) or _CURRENCY_PATTERN.match(str(c).strip()))
    return numeric_count < len(non_empty) * 0.5


def _build_xlsx(
    all_page_tables: Dict[int, List[List[List[str]]]],
    total_pages: int,
    output_path: str,
) -> Dict[str, Any]:
    """
    Build an enterprise-grade XLSX workbook from extracted table data.
    Features:
      - Header styling: Slate fill (#F1F5F9), bold dark typography (#0F172A).
      - Grid lines and clean cell borders (#CBD5E1).
      - Smart alignments: Right-aligned numbers & currencies, centered dates, left-aligned text.
      - Consolidated master worksheet for multi-page documents with consistent schemas.
      - Individual page worksheets (Page 1, Page 2...).
      - Column width auto-sizing with text-wrapping support.
    """
    wb = openpyxl.Workbook()
    # Remove default empty sheet
    default_sheet = wb.active
    if default_sheet:
        wb.remove(default_sheet)

    total_tables = 0
    total_rows_written = 0
    pages_with_data = 0

    # Styling definitions
    header_font = Font(name='Calibri', size=11, bold=True, color='0F172A')
    header_fill = PatternFill(start_color='F1F5F9', end_color='F1F5F9', fill_type='solid')
    header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    normal_font = Font(name='Calibri', size=11, color='1E293B')
    num_alignment = Alignment(horizontal='right', vertical='center')
    date_alignment = Alignment(horizontal='center', vertical='center')
    text_alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    def write_table_to_sheet(
        ws: openpyxl.worksheet.worksheet.Worksheet,
        table: List[List[str]],
        start_row: int
    ) -> int:
        nonlocal total_rows_written
        curr_row = start_row
        if not table:
            return curr_row

        num_rows = len(table)
        max_cols = max(len(r) for r in table) if table else 0
        if max_cols == 0:
            return curr_row

        norm_table = [list(r) + [''] * (max_cols - len(r)) for r in table]

        # 1. Detect Merged Cell Ranges (Colspan & Rowspan)
        merges_to_apply: List[Tuple[int, int, int, int]] = []
        merged_covered = set()

        # A) Horizontal Spans in header rows
        for r_idx in range(min(num_rows, 3)):
            row_data = norm_table[r_idx]
            skip_c = -1
            for c_idx in range(max_cols):
                if c_idx <= skip_c:
                    continue
                val = str(row_data[c_idx] or "").strip()
                if val:
                    consec_empty = 0
                    lookahead = c_idx + 1
                    while lookahead < max_cols and not str(norm_table[r_idx][lookahead] or "").strip():
                        consec_empty += 1
                        lookahead += 1

                    if consec_empty > 0 and r_idx + 1 < num_rows:
                        child_active = sum(
                            1 for chk_c in range(c_idx, c_idx + consec_empty + 1)
                            if str(norm_table[r_idx + 1][chk_c] or "").strip()
                        )
                        if child_active >= 2:
                            r_start = curr_row + r_idx
                            c_start = c_idx + 1
                            r_end = curr_row + r_idx
                            c_end = c_idx + consec_empty + 1
                            merges_to_apply.append((r_start, c_start, r_end, c_end))
                            for mc in range(c_start + 1, c_end + 1):
                                merged_covered.add((r_start, mc))
                            skip_c = c_idx + consec_empty

        # B) Vertical Spans (e.g. Day spanning row 0 and row 1)
        if num_rows >= 2:
            for c_idx in range(max_cols):
                val_r0 = str(norm_table[0][c_idx] or "").strip()
                val_r1 = str(norm_table[1][c_idx] or "").strip()
                if val_r0 and not val_r1:
                    has_sibling = any(str(norm_table[1][sc] or "").strip() for sc in range(max_cols) if sc != c_idx)
                    if has_sibling:
                        r_start = curr_row
                        c_start = c_idx + 1
                        r_end = curr_row + 1
                        c_end = c_idx + 1
                        merges_to_apply.append((r_start, c_start, r_end, c_end))
                        merged_covered.add((r_end, c_start))

        # 2. Populate cells
        has_header = _is_likely_header(norm_table[0], len(norm_table))

        for r_idx, row_data in enumerate(norm_table):
            excel_r = curr_row + r_idx
            is_hdr_row = (r_idx == 0 and has_header) or (r_idx == 1 and has_header and len(norm_table) > 2)

            for col_idx, cell_value in enumerate(row_data):
                excel_c = col_idx + 1
                cell = ws.cell(row=excel_r, column=excel_c)
                cell.border = thin_border

                if is_hdr_row:
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = header_alignment
                else:
                    cell.font = normal_font

                typed_val, num_fmt = detect_cell_type(cell_value)
                if num_fmt:
                    cell.number_format = num_fmt

                if (excel_r, excel_c) not in merged_covered:
                    cell.value = typed_val
                    if not is_hdr_row:
                        if num_fmt and ('$' in num_fmt or '€' in num_fmt or '£' in num_fmt or '¥' in num_fmt or '0.00' in num_fmt or '%' in num_fmt):
                            cell.alignment = num_alignment
                        elif isinstance(typed_val, (int, float)):
                            cell.alignment = num_alignment
                        elif isinstance(typed_val, datetime) or (num_fmt and 'YYYY' in num_fmt):
                            cell.alignment = date_alignment
                        else:
                            cell.alignment = text_alignment
                else:
                    cell.value = None

            total_rows_written += 1

        # 3. Format complete borders and fill across all merged rectangles
        for r_start, c_start, r_end, c_end in merges_to_apply:
            for mr in range(r_start, r_end + 1):
                for mc in range(c_start, c_end + 1):
                    cell_in_box = ws.cell(row=mr, column=mc)
                    cell_in_box.border = thin_border
                    cell_in_box.fill = header_fill

        # 4. Apply openpyxl ws.merge_cells
        for r_start, c_start, r_end, c_end in merges_to_apply:
            ws.merge_cells(
                start_row=r_start,
                start_column=c_start,
                end_row=r_end,
                end_column=c_end
            )
            top_left = ws.cell(row=r_start, column=c_start)
            top_left.alignment = header_alignment

        return curr_row + num_rows


    def auto_fit_columns(ws: openpyxl.worksheet.worksheet.Worksheet):
        for col_idx in range(1, (ws.max_column or 0) + 1):
            col_letter = get_column_letter(col_idx)
            max_len = 10  # Minimum comfortable column width
            for row_idx in range(1, (ws.max_row or 0) + 1):
                c = ws.cell(row=row_idx, column=col_idx)
                if c.value is not None:
                    line_lens = [len(str(line)) for line in str(c.value).split('\n')]
                    cell_max = max(line_lens) if line_lens else 0
                    max_len = max(max_len, min(cell_max + 3, 60))
            ws.column_dimensions[col_letter].width = max_len

    # Multi-page consolidation check: identify primary data tables (with >= 2 columns) per page
    valid_pages = [p for p in range(total_pages) if all_page_tables.get(p)]
    can_consolidate = len(valid_pages) >= 2

    consolidated_rows: List[List[str]] = []
    if can_consolidate:
        primary_tables: Dict[int, List[List[str]]] = {}
        first_cols = None

        for p_idx in valid_pages:
            tables = all_page_tables[p_idx]
            # Select table with largest grid volume having at least 2 columns
            multi_col_tables = [
                _normalize_table_columns(t) for t in tables
                if t and len(_normalize_table_columns(t)[0]) >= 2
            ]

            if not multi_col_tables:
                can_consolidate = False
                break

            chosen_table = max(multi_col_tables, key=lambda t: len(t) * len(t[0]))
            num_cols = len(chosen_table[0])

            if first_cols is None:
                first_cols = num_cols
            elif first_cols != num_cols:
                can_consolidate = False
                break

            primary_tables[p_idx] = chosen_table

        if can_consolidate and first_cols and len(primary_tables) == len(valid_pages):
            for idx, p_idx in enumerate(valid_pages):
                norm_t = primary_tables[p_idx]
                has_hdr = _is_likely_header(norm_t[0], len(norm_t))
                if idx == 0:
                    consolidated_rows.extend(norm_t)
                else:
                    rows_to_add = norm_t[1:] if has_hdr else norm_t
                    consolidated_rows.extend(rows_to_add)

            if len(consolidated_rows) > 1:
                ws_cons = wb.create_sheet(title="All Data (Consolidated)")
                write_table_to_sheet(ws_cons, consolidated_rows, 1)
                auto_fit_columns(ws_cons)

    # Individual page worksheets
    for page_idx in range(total_pages):
        page_tables = all_page_tables.get(page_idx, [])
        if not page_tables:
            continue

        pages_with_data += 1
        ws_title = f"Page {page_idx + 1}"
        ws = wb.create_sheet(title=ws_title[:31])

        current_row = 1
        for table_idx, table in enumerate(page_tables):
            if not table:
                continue

            table = _normalize_table_columns(table)
            total_tables += 1

            if table_idx > 0:
                current_row += 2  # Visual separation between tables

            current_row = write_table_to_sheet(ws, table, current_row)

        auto_fit_columns(ws)

    # Empty document fallback sheet
    if not wb.sheetnames:
        ws = wb.create_sheet(title="Results")
        ws.cell(row=1, column=1, value="No tables or structured data found in the PDF document.")
        ws.cell(row=2, column=1, value="The PDF may contain only images, raw scanned graphics, or non-tabular content.")
        ws.column_dimensions['A'].width = 65

    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    wb.save(output_path)


    return {
        "total_tables_extracted": total_tables,
        "total_rows": total_rows_written,
        "pages_with_data": pages_with_data,
        "total_pages_processed": total_pages,
        "worksheets_created": len(wb.sheetnames),
    }


# ---------------------------------------------------------------------------
# Main Converter Class
# ---------------------------------------------------------------------------

class PdfToExcelConverter(BaseConverter):
    """
    Production-grade PDF to Excel converter.
    Extracts tables, financial data, invoices, and bank statements
    from PDF documents into structured XLSX spreadsheets.
    """

    @property
    def tool_id(self) -> str:
        return "pdf-to-excel"

    @property
    def name(self) -> str:
        return "PDF to Excel"

    @property
    def supported_inputs(self) -> List[str]:
        return ["pdf"]

    @property
    def output_extension(self) -> str:
        return "xlsx"

    @property
    def output_mime_type(self) -> str:
        return "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    def validate_inputs(self, input_paths: List[str], options: Dict[str, Any]) -> None:
        super().validate_inputs(input_paths, options)
        src_path = input_paths[0]

        # 1. File existence & size validation
        if not os.path.exists(src_path):
            raise FileValidationError(f"File does not exist: {src_path}")
        size_bytes = os.path.getsize(src_path)
        if size_bytes == 0:
            raise FileValidationError("Uploaded PDF file is empty (0 bytes).")
        if size_bytes > 100 * 1024 * 1024:
            raise FileValidationError("File size exceeds the 100 MB upload limit.")

        # 2. Magic bytes validation: must contain %PDF-
        with open(src_path, "rb") as f:
            header = f.read(1024)
        if b"%PDF-" not in header:
            raise FileValidationError("Invalid file format: Not a genuine PDF document.")

        # 3. Document integrity, password & bomb checks
        try:
            doc = fitz.open(src_path)
            if doc.is_encrypted:
                is_authenticated = False
                pwd = options.get("password", "")
                try:
                    is_authenticated = bool(doc.authenticate(pwd))
                except Exception:
                    pass
                if not is_authenticated:
                    doc.close()
                    raise FileValidationError(
                        "This PDF is password-protected. Please provide the correct password or unlock the PDF first."
                    )
            if doc.page_count < 1:
                doc.close()
                raise FileValidationError("PDF contains 0 pages and is invalid.")
            if doc.page_count > 1000:
                doc.close()
                raise FileValidationError("Document exceeds maximum supported limit of 1,000 pages.")

            # Decompression bomb / huge dimension protection (max 28800 pt = 400 inches)
            for p in doc:
                r = p.rect
                if r.width > 28800 or r.height > 28800:
                    doc.close()
                    raise FileValidationError("PDF page dimensions exceed safety limits (potential decompression attack).")
            doc.close()
        except FileValidationError:
            raise
        except Exception as e:
            raise FileValidationError(f"File is corrupted or not a valid PDF: {str(e)}")

    async def convert(
        self,
        input_paths: List[str],
        output_dir: str,
        options: Dict[str, Any]
    ) -> ConversionResult:
        self.validate_inputs(input_paths, options)
        os.makedirs(output_dir, exist_ok=True)
        src_path = input_paths[0]
        out_filename = f"converted_{uuid.uuid4().hex[:8]}.xlsx"
        out_path = os.path.join(output_dir, out_filename)



        # ---------------------------------------------------------------
        # Stage 0: Enterprise High-Accuracy Document Processing Engine
        # ---------------------------------------------------------------
        try:
            logger.info("Executing EnterpriseTableExtractor engine...")
            extractor = EnterpriseTableExtractor()
            meta = await asyncio.to_thread(extractor.process_pdf_document, src_path, out_path, options)

            # If tables were detected and written, validate and return
            if meta.get("total_tables_extracted", 0) > 0:
                self.validate_output(out_path)
                return ConversionResult(
                    output_path=out_path,
                    output_filename=out_filename,
                    mime_type=self.output_mime_type,
                    size_bytes=os.path.getsize(out_path),
                    metadata={
                        "page_count": meta.get("total_pages", 1),
                        "extraction_method": "enterprise_table_extractor",
                        "is_scanned": meta.get("is_scanned", False),
                        **meta,
                    }
                )
        except Exception as exc:
            logger.warning(
                f"EnterpriseTableExtractor encountered an issue, proceeding to secondary fallback pipeline: {exc}",
                exc_info=True
            )

        force_ocr = options.get("ocr", False) or options.get("force_ocr", False)

        # Open and inspect the PDF
        doc = open_and_validate_pdf(src_path)
        total_pages = doc.page_count
        page_indices = list(range(total_pages))

        # Detect if PDF is scanned (low text density)
        page_texts = [page.get_text().strip() for page in doc]
        total_chars = sum(len(t) for t in page_texts)
        has_native_text = total_chars >= max(30, total_pages * 15)
        is_scanned = not has_native_text
        if force_ocr and has_native_text:
            logger.info("OCR requested, but PDF contains native selectable text. Automatically skipping OCR to preserve speed and vector quality.")

        logger.info(
            f"PDF to Excel fallback: {total_pages} pages, {total_chars} chars, "
            f"scanned={'yes' if is_scanned else 'no'}, force_ocr={force_ocr}"
        )

        # ---------------------------------------------------------------
        # Stage 1: Extract tables
        # ---------------------------------------------------------------
        all_page_tables: Dict[int, List[List[List[str]]]] = {}

        if not is_scanned:
            # Primary: pdfplumber extraction
            logger.info("Running pdfplumber table extraction...")
            all_page_tables = _extract_tables_pdfplumber(src_path, page_indices)
            extraction_method = "pdfplumber"

            # Check coverage - if pdfplumber missed pages that have text, try tabula
            missing_pages = [
                idx for idx in page_indices
                if idx not in all_page_tables and len(page_texts[idx]) > 20
            ]

            if missing_pages:
                logger.info(f"Trying tabula-py for {len(missing_pages)} pages with no pdfplumber results...")
                tabula_results = _extract_tables_tabula(src_path, missing_pages)
                for page_idx, tables in tabula_results.items():
                    if page_idx not in all_page_tables:
                        all_page_tables[page_idx] = tables
                        extraction_method = "pdfplumber+tabula"

            # Final fallback: for pages with text but no extracted tables,
            # do a raw text grid extraction using PyMuPDF
            still_missing = [
                idx for idx in page_indices
                if idx not in all_page_tables and len(page_texts[idx]) > 20
            ]
            if still_missing:
                logger.info(f"Raw text grid fallback for {len(still_missing)} remaining pages...")
                for page_idx in still_missing:
                    text = page_texts[page_idx]
                    lines = [line.strip() for line in text.split('\n') if line.strip()]
                    if lines:
                        grid = []
                        for line in lines:
                            cells = re.split(r'\t|  {2,}', line)
                            cells = [c.strip() for c in cells]
                            if any(c for c in cells):
                                grid.append(cells)
                        if grid:
                            all_page_tables[page_idx] = [grid]
                            extraction_method = "pdfplumber+text"

        else:
            # Scanned PDF -> OCR pipeline
            logger.info("Scanned PDF detected - running OCR extraction pipeline...")
            extraction_method = "ocr"

            # Find Tesseract
            from app.services.engine.office import find_tesseract_bin
            tesseract_bin = find_tesseract_bin()

            if tesseract_bin:
                all_page_tables = _extract_tables_ocr(doc, page_indices, tesseract_bin)
            else:
                logger.warning("Tesseract not available. Attempting raw text extraction from scanned PDF...")
                # Even scanned PDFs sometimes have a thin text layer
                for page_idx in page_indices:
                    text = page_texts[page_idx] if page_idx < len(page_texts) else ""
                    if text and len(text.strip()) > 5:
                        lines = [line.strip() for line in text.split('\n') if line.strip()]
                        if lines:
                            grid = []
                            for line in lines:
                                cells = re.split(r'\t|  {2,}', line)
                                cells = [c.strip() for c in cells]
                                if any(c for c in cells):
                                    grid.append(cells)
                            if grid:
                                all_page_tables[page_idx] = [grid]
                extraction_method = "text_fallback"

        doc.close()

        # ---------------------------------------------------------------
        # Stage 2: Generate XLSX
        # ---------------------------------------------------------------
        out_filename = f"converted_{uuid.uuid4().hex[:8]}.xlsx"
        out_path = os.path.join(output_dir, out_filename)

        logger.info(f"Generating XLSX with {len(all_page_tables)} pages of data...")
        generation_meta = _build_xlsx(all_page_tables, total_pages, out_path)

        self.validate_output(out_path)

        return ConversionResult(
            output_path=out_path,
            output_filename=out_filename,
            mime_type=self.output_mime_type,
            size_bytes=os.path.getsize(out_path),
            metadata={
                "page_count": total_pages,
                "extraction_method": extraction_method,
                "is_scanned": is_scanned,
                **generation_meta,
            }
        )
