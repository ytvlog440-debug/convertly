"""
Convertly V2 — Dedicated Merged-Cell Reconstruction Test Suite
Validates dynamic openpyxl ws.merge_cells() generation for:
  - Multi-column headers ('Morning' -> B2:F2, 'Replica' -> H2:K2)
  - Vertical multi-level headers ('Day' -> A2:A3, 'Break' -> G2:G3)
  - Centered & left-aligned header text block dimensions
  - Preserving row spans and column spans
  - Both EnterpriseTableExtractor and fallback XLSX pipelines
"""

import os
import openpyxl
import pytest
import fitz

from app.services.engine.table_extractor.models import (
    TableCell,
    TableRow,
    TableBlock,
    PageLayout,
)
from app.services.engine.table_extractor.reconstruction.reconstructor import CellReconstructor
from app.services.engine.table_extractor.excel_writer import ExcelDocumentWriter
from app.services.engine.table_extractor.excel.writer import EnterpriseExcelWriter
from app.services.engine.pdf_to_excel import PdfToExcelConverter, _build_xlsx


TEST_TMP = os.path.join(os.path.dirname(__file__), "tmp_merged_tests")


def setup_module():
    os.makedirs(TEST_TMP, exist_ok=True)


def teardown_module():
    import shutil
    shutil.rmtree(TEST_TMP, ignore_errors=True)


def test_reconstruct_morning_replica_timetable_headers():
    """
    Verify multi-level timetable header reconstruction:
      - 'Morning' spans B2:F2 (cols 2..6)
      - 'Replica' spans H2:K2 (cols 8..11)
      - 'Day' spans A2:A3 (col 1, rows 2..3)
      - 'Break' spans G2:G3 (col 7, rows 2..3)
    """
    rec = CellReconstructor()

    # Row 0: Section headers (Row 2 in Excel after Title)
    r0_cells = [
        TableCell("Day", (50, 100, 110, 120), 0, 0, is_header=True),
        TableCell("Morning", (110, 100, 170, 120), 0, 1, is_header=True),
        TableCell("", (170, 100, 230, 120), 0, 2, is_header=True),
        TableCell("", (230, 100, 290, 120), 0, 3, is_header=True),
        TableCell("", (290, 100, 350, 120), 0, 4, is_header=True),
        TableCell("", (350, 100, 410, 120), 0, 5, is_header=True),
        TableCell("Break", (410, 100, 470, 120), 0, 6, is_header=True),
        TableCell("Replica", (470, 100, 530, 120), 0, 7, is_header=True),
        TableCell("", (530, 100, 590, 120), 0, 8, is_header=True),
        TableCell("", (590, 100, 650, 120), 0, 9, is_header=True),
        TableCell("", (650, 100, 710, 120), 0, 10, is_header=True),
    ]
    r0 = TableRow(r0_cells, 0, (50, 100, 710, 120), is_header=True)

    # Row 1: Time slot granular subheaders (Row 3 in Excel)
    slots = ["Day", "09:00", "10:00", "11:00", "12:00", "13:00", "Break", "14:00", "15:00", "16:00", "17:00"]
    r1_cells = [
        TableCell(slots[i] if i not in (0, 6) else "", (50 + i * 60, 120, 110 + i * 60, 140), 1, i, is_header=True)
        for i in range(11)
    ]
    r1 = TableRow(r1_cells, 1, (50, 120, 710, 140), is_header=True)

    # Row 2: Data row (Monday)
    r2_cells = [TableCell(f"Data_{i}", (50 + i * 60, 140, 110 + i * 60, 160), 2, i) for i in range(11)]
    r2 = TableRow(r2_cells, 2, (50, 140, 710, 160))

    tbl = TableBlock((50, 100, 710, 160), [r0, r1, r2])
    repaired = rec.detect_merged_spans(tbl)

    assert repaired.has_merged_cells is True
    # Morning spans cols 1..5 (5 columns)
    assert repaired.rows[0].cells[1].colspan == 5
    assert repaired.rows[0].cells[1].text == "Morning"
    # Replica spans cols 7..10 (4 columns)
    assert repaired.rows[0].cells[7].colspan == 4
    assert repaired.rows[0].cells[7].text == "Replica"
    # Day spans rows 0..1 (2 rows)
    assert repaired.rows[0].cells[0].rowspan == 2
    # Break spans rows 0..1 (2 rows)
    assert repaired.rows[0].cells[6].rowspan == 2

    # Render into Excel worksheet with Title at row 1
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Weekly Department Lecture Timetable"
    writer = ExcelDocumentWriter()
    writer._render_table_to_sheet(ws, repaired.rows, start_row=2)

    merged_ranges = [str(r) for r in ws.merged_cells.ranges]
    # Check exact expected Excel coordinates:
    assert "B2:F2" in merged_ranges
    assert "H2:K2" in merged_ranges
    assert "A2:A3" in merged_ranges
    assert "G2:G3" in merged_ranges

    assert ws["B2"].value == "Morning"
    assert ws["H2"].value == "Replica"
    assert ws["A2"].value == "Day"
    assert ws["G2"].value == "Break"
    wb.close()


def test_reconstruct_centered_headers():
    """
    Verify that if 'Morning' is centered at Col 3 (D) and 'Replica' is centered at Col 8 (I),
    they are correctly relocated to their top-left cells (B2 and H2) with proper colspans.
    """
    rec = CellReconstructor()

    # Row 0 with centered text
    r0_cells = [
        TableCell("Day", (50, 100, 110, 120), 0, 0, is_header=True),
        TableCell("", (110, 100, 170, 120), 0, 1, is_header=True),
        TableCell("", (170, 100, 230, 120), 0, 2, is_header=True),
        TableCell("Morning", (230, 100, 290, 120), 0, 3, is_header=True),  # Centered at Col 3
        TableCell("", (290, 100, 350, 120), 0, 4, is_header=True),
        TableCell("", (350, 100, 410, 120), 0, 5, is_header=True),
        TableCell("Break", (410, 100, 470, 120), 0, 6, is_header=True),
        TableCell("", (470, 100, 530, 120), 0, 7, is_header=True),
        TableCell("Replica", (530, 100, 590, 120), 0, 8, is_header=True),  # Centered at Col 8
        TableCell("", (590, 100, 650, 120), 0, 9, is_header=True),
        TableCell("", (650, 100, 710, 120), 0, 10, is_header=True),
    ]
    r0 = TableRow(r0_cells, 0, (50, 100, 710, 120), is_header=True)

    slots = ["Day", "09:00", "10:00", "11:00", "12:00", "13:00", "Break", "14:00", "15:00", "16:00", "17:00"]
    r1_cells = [
        TableCell(slots[i] if i not in (0, 6) else "", (50 + i * 60, 120, 110 + i * 60, 140), 1, i, is_header=True)
        for i in range(11)
    ]
    r1 = TableRow(r1_cells, 1, (50, 120, 710, 140), is_header=True)
    r2_cells = [TableCell(f"Data_{i}", (50 + i * 60, 140, 110 + i * 60, 160), 2, i) for i in range(11)]
    r2 = TableRow(r2_cells, 2, (50, 140, 710, 160))

    tbl = TableBlock((50, 100, 710, 160), [r0, r1, r2])
    repaired = rec.detect_merged_spans(tbl)

    # Morning should be relocated to index 1 with colspan 5
    assert repaired.rows[0].cells[1].colspan == 5
    assert repaired.rows[0].cells[1].text == "Morning"
    # Replica should be relocated to index 7 with colspan 4
    assert repaired.rows[0].cells[7].colspan == 4
    assert repaired.rows[0].cells[7].text == "Replica"


def test_fallback_build_xlsx_merged_cells():
    """
    Verify that the fallback _build_xlsx / write_table_to_sheet pipeline also
    detects multi-column and vertical spans and calls ws.merge_cells().
    """
    out_xlsx = os.path.join(TEST_TMP, "fallback_merged.xlsx")
    table_grid = [
        ["Day", "Morning", "", "", "", "", "Break", "Replica", "", "", ""],
        ["", "09:00", "10:00", "11:00", "12:00", "13:00", "", "14:00", "15:00", "16:00", "17:00"],
        ["Monday", "Math", "Physics", "Chemistry", "Biology", "Lab", "Lunch", "CS", "History", "Music", "Art"],
    ]

    _build_xlsx({0: [table_grid]}, total_pages=1, output_path=out_xlsx)

    wb = openpyxl.load_workbook(out_xlsx)
    ws = wb.active
    ranges = [str(r) for r in ws.merged_cells.ranges]
    assert "B1:F1" in ranges
    assert "H1:K1" in ranges
    assert "A1:A2" in ranges
    assert "G1:G2" in ranges
    assert ws["B1"].value == "Morning"
    assert ws["H1"].value == "Replica"
    assert ws["A1"].value == "Day"
    assert ws["G1"].value == "Break"
    wb.close()


@pytest.mark.asyncio
async def test_end_to_end_pdf_to_excel_merged_timetable():
    """
    Create an actual vector PDF with Morning spanning B2:F2 and Replica spanning H2:K2,
    convert using PdfToExcelConverter, and verify the resulting XLSX.
    """
    pdf_path = os.path.join(TEST_TMP, "e2e_merged_timetable.pdf")
    doc = fitz.open()
    page = doc.new_page(width=792, height=612)

    # Title
    page.insert_text((50, 40), "Conference Master Schedule", fontsize=14)

    # Draw lines
    y_lines = [70, 95, 120, 145]
    for y in y_lines:
        page.draw_line((50, y), (710, y))

    x_coords = [50, 110, 170, 230, 290, 350, 410, 470, 530, 590, 650, 710]
    # Vertical lines for outer boundaries
    page.draw_line((50, 70), (50, 145))
    page.draw_line((710, 70), (710, 145))
    # Divider for Break
    page.draw_line((410, 70), (410, 145))
    page.draw_line((470, 70), (470, 145))

    # In Row 0: Header text
    page.insert_text((55, 87), "Day", fontsize=10)
    page.insert_text((200, 87), "Morning", fontsize=10)
    page.insert_text((420, 87), "Break", fontsize=10)
    page.insert_text((550, 87), "Replica", fontsize=10)

    # In Row 1: Subheaders
    slots = ["09:00", "10:00", "11:00", "12:00", "13:00", "14:00", "15:00", "16:00", "17:00"]
    slot_cols = [1, 2, 3, 4, 5, 7, 8, 9, 10]
    for s_idx, col_i in enumerate(slot_cols):
        page.insert_text((55 + col_i * 60, 112), slots[s_idx], fontsize=9)

    # In Row 2: Data
    page.insert_text((55, 137), "Monday", fontsize=9)
    for col_i in range(1, 11):
        if col_i != 6:
            page.insert_text((55 + col_i * 60, 137), f"Session {col_i}", fontsize=9)

    doc.save(pdf_path)
    doc.close()

    converter = PdfToExcelConverter()
    out_dir = os.path.join(TEST_TMP, "e2e_out")
    res = await converter.convert([pdf_path], out_dir, {"ocr": False})

    assert os.path.exists(res.output_path)
    wb = openpyxl.load_workbook(res.output_path)
    ws = wb.active
    assert ws is not None
    merged_ranges = [str(r) for r in ws.merged_cells.ranges]
    assert len(merged_ranges) >= 1
    wb.close()
