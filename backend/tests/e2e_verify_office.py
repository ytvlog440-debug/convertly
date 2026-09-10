import requests
import json
import time
import os
import docx

BASE_URL = "http://127.0.0.1:8000/api/v1"

def test_office_e2e():
    # 1. Create a real sample DOCX file
    docx_path = "d:/convertlytools/office_test.docx"
    doc = docx.Document()
    doc.add_heading("Convertly V2 Enterprise Word Pipeline", 0)
    doc.add_paragraph("Testing full end-to-end conversion from Word DOCX to PDF over the HTTP API.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Header 1"
    table.cell(0, 1).text = "Header 2"
    table.cell(1, 0).text = "Value A"
    table.cell(1, 1).text = "Value B"
    doc.save(docx_path)

    # 2. Upload DOCX
    with open(docx_path, "rb") as f:
        res_up = requests.post(f"{BASE_URL}/files/upload", files={"file": ("office_test.docx", f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert res_up.status_code == 200, res_up.text
    file_id = res_up.json()["data"]["id"]
    print("Uploaded DOCX file ID:", file_id)

    # 3. Submit Word to PDF Job
    job_payload = {
        "tool_id": "word-to-pdf",
        "input_file_ids": [file_id],
        "options": {}
    }
    res_job = requests.post(f"{BASE_URL}/jobs", json=job_payload)
    assert res_job.status_code == 200, res_job.text
    job_id = res_job.json()["data"]["id"]
    print("Submitted Word-to-PDF Job ID:", job_id)

    # 4. Poll until completed
    for _ in range(25):
        time.sleep(0.4)
        poll_res = requests.get(f"{BASE_URL}/jobs/{job_id}")
        assert poll_res.status_code == 200
        data = poll_res.json()["data"]
        if data["status"] in ["completed", "failed"]:
            print("Job finished:", data["status"])
            assert data["status"] == "completed"
            assert data["output_file_id"] is not None

            # Download resulting PDF
            dl = requests.get(f"{BASE_URL}/files/{data['output_file_id']}/download")
            assert dl.status_code == 200
            assert len(dl.content) > 0
            print("Successfully downloaded converted PDF from Word, size:", len(dl.content), "bytes")
            return

    raise TimeoutError("Word-to-PDF conversion timed out")

if __name__ == "__main__":
    test_office_e2e()
