import os
import shutil
import pytest
import zipfile
import fitz  # PyMuPDF
from app.services.engine.office import (
    WordToPdfConverter,
    PdfToWordConverter,
    ExcelToPdfConverter,
    PptToPdfConverter
)
from app.core.errors import FileValidationError

TEST_TMP_DIR = "./temp_office_tests"


@pytest.fixture(autouse=True)
def setup_teardown():
    os.makedirs(TEST_TMP_DIR, exist_ok=True)
    yield
    if os.path.exists(TEST_TMP_DIR):
        shutil.rmtree(TEST_TMP_DIR, ignore_errors=True)


def create_sample_docx(path: str) -> str:
    import docx
    doc = docx.Document()
    doc.add_heading("Convertly V2 Enterprise Word Converter", 0)
    doc.add_paragraph("This is a high-fidelity test paragraph demonstrating Word to PDF conversion.")
    
    table = doc.add_table(rows=3, cols=3)
    for r in range(3):
        for c in range(3):
            table.cell(r, c).text = f"Cell R{r+1}C{c+1}"
    doc.save(path)
    return path


def create_sample_xlsx(path: str) -> str:
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Financial Summary"
    ws.append(["Quarter", "Revenue ($M)", "Growth (%)"])
    ws.append(["Q1", 12.5, "15%"])
    ws.append(["Q2", 15.2, "22%"])
    ws.append(["Q3", 18.9, "24%"])
    wb.save(path)
    return path


def create_sample_pptx(path: str) -> str:
    from pptx import Presentation
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Convertly V2 Presentation"
    slide.placeholders[1].text = "Slide 1 Subtitle and description text"
    
    slide2 = prs.slides.add_slide(prs.slide_layouts[1])
    slide2.shapes.title.text = "Slide 2 Architecture Overview"
    slide2.placeholders[1].text = "Production-grade native slide renderer"
    prs.save(path)
    return path


@pytest.mark.asyncio
async def test_word_to_pdf_conversion():
    docx_path = os.path.join(TEST_TMP_DIR, "test_document.docx")
    create_sample_docx(docx_path)

    converter = WordToPdfConverter()
    result = await converter.convert([docx_path], TEST_TMP_DIR, {})

    assert os.path.exists(result.output_path)
    assert result.mime_type == "application/pdf"
    assert result.size_bytes > 0

    # Verify output is a readable PDF
    doc = fitz.open(result.output_path)
    assert doc.page_count >= 1
    doc.close()


@pytest.mark.asyncio
async def test_pdf_to_word_conversion():
    pdf_path = os.path.join(TEST_TMP_DIR, "test_pdf_to_docx.pdf")
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 100), "Convertly V2 PDF to Word Reconstruction Test", fontsize=18)
    doc.save(pdf_path)
    doc.close()

    converter = PdfToWordConverter()
    result = await converter.convert([pdf_path], TEST_TMP_DIR, {})

    assert os.path.exists(result.output_path)
    assert result.mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    assert result.size_bytes > 0

    # Verify output is a valid DOCX package
    assert zipfile.is_zipfile(result.output_path)
    with zipfile.ZipFile(result.output_path, "r") as z:
        assert "word/document.xml" in z.namelist()


@pytest.mark.asyncio
async def test_excel_to_pdf_conversion():
    xlsx_path = os.path.join(TEST_TMP_DIR, "test_spreadsheet.xlsx")
    create_sample_xlsx(xlsx_path)

    converter = ExcelToPdfConverter()
    result = await converter.convert([xlsx_path], TEST_TMP_DIR, {})

    assert os.path.exists(result.output_path)
    assert result.mime_type == "application/pdf"
    assert result.size_bytes > 0

    doc = fitz.open(result.output_path)
    assert doc.page_count >= 1
    doc.close()


@pytest.mark.asyncio
async def test_pptx_to_pdf_conversion():
    pptx_path = os.path.join(TEST_TMP_DIR, "test_presentation.pptx")
    create_sample_pptx(pptx_path)

    converter = PptToPdfConverter()
    result = await converter.convert([pptx_path], TEST_TMP_DIR, {})

    assert os.path.exists(result.output_path)
    assert result.mime_type == "application/pdf"
    assert result.size_bytes > 0

    doc = fitz.open(result.output_path)
    assert doc.page_count == 2
    doc.close()


@pytest.mark.asyncio
async def test_corrupted_office_file_rejected():
    corrupt_path = os.path.join(TEST_TMP_DIR, "corrupt.docx")
    with open(corrupt_path, "wb") as f:
        f.write(b"NOT_A_REAL_ZIP_FILE_DATA_12345")

    converter = WordToPdfConverter()
    with pytest.raises(FileValidationError, match="not a valid OpenXML document"):
        await converter.convert([corrupt_path], TEST_TMP_DIR, {})
