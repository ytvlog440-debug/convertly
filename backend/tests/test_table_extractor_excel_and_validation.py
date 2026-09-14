"""
Convertly V2 — Table Validation & High-Fidelity Excel Writer Test Suite
Validates TableValidator arithmetic and structural audits, and EnterpriseExcelWriter
OpenXML XLSX generation with typed values, number formatting, indentation, and merged cells.
"""

import io
import os
import openpyxl
import pytest

from app.services.engine.table_extractor.models import (
    TableCell,
    TableRow,
    TableBlock,
)
from app.services.engine.table_extractor.validation.validator import TableValidator
from app.services.engine.table_extractor.excel.writer import EnterpriseExcelWriter
from app.services.engine.table_extractor.container import get_default_container


# =====================================================================
# 1. Table Validator Tests
# =====================================================================

def test_table_validator_math_reconciliation_pass():
    validator = TableValidator()

    # Header row
    h1 = TableCell("Item", (50, 100, 150, 120), 0, 0, is_header=True)
    h2 = TableCell("Amount", (150, 100, 250, 120), 0, 1, is_header=True)
    r_hdr = TableRow([h1, h2], row_idx=0, is_header=True)

    # Data row 1: $100.00
    c1 = TableCell("Service Fee", (50, 125, 150, 140), 1, 0)
    c2 = TableCell("$100.00", (150, 125, 250, 140), 1, 1, data_type="currency", typed_value=100.0)
    r1 = TableRow([c1, c2], row_idx=1)

    # Data row 2: $250.50
    c3 = TableCell("Server Upgrade", (50, 145, 150, 160), 2, 0)
    c4 = TableCell("$250.50", (150, 145, 250, 160), 2, 1, data_type="currency", typed_value=250.50)
    r2 = TableRow([c3, c4], row_idx=2)

    # Total row: Total $350.50
    t1 = TableCell("Total Due", (50, 165, 150, 180), 3, 0)
    t2 = TableCell("$350.50", (150, 165, 250, 180), 3, 1, data_type="currency", typed_value=350.50)
    r_tot = TableRow([t1, t2], row_idx=3)

    table = TableBlock(bbox=(50, 100, 250, 180), rows=[r_hdr, r1, r2, r_tot])

    report = validator.validate(table)
    assert report.is_valid is True
    assert report.math_checks_passed == 1
    assert report.math_checks_failed == 0


def test_table_validator_math_reconciliation_mismatch():
    validator = TableValidator()

    h1 = TableCell("Item", (50, 100, 150, 120), 0, 0, is_header=True)
    h2 = TableCell("Amount", (150, 100, 250, 120), 0, 1, is_header=True)
    r_hdr = TableRow([h1, h2], row_idx=0, is_header=True)

    c1 = TableCell("Item A", (50, 125, 150, 140), 1, 0)
    c2 = TableCell("$50.00", (150, 125, 250, 140), 1, 1, data_type="currency", typed_value=50.0)
    r1 = TableRow([c1, c2], row_idx=1)

    # Total reported as $999.00 (mismatch with $50.00)
    t1 = TableCell("Subtotal", (50, 145, 150, 160), 2, 0)
    t2 = TableCell("$999.00", (150, 145, 250, 160), 2, 1, data_type="currency", typed_value=999.0)
    r_tot = TableRow([t1, t2], row_idx=2)

    table = TableBlock(bbox=(50, 100, 250, 160), rows=[r_hdr, r1, r_tot])

    report = validator.validate(table)
    assert report.is_valid is False
    assert report.math_checks_failed == 1
    assert len(report.violations) == 1
    assert report.violations[0]["type"] == "math_mismatch"


# =====================================================================
# 2. Enterprise Excel Writer Tests
# =====================================================================

def test_excel_writer_generation_and_typing(tmp_path):
    writer = EnterpriseExcelWriter()

    # Create table with headers, currency, numbers, percentages, dates, and indentation
    h1 = TableCell("Financial Line Item", (50, 50, 250, 70), 0, 0, is_header=True)
    h2 = TableCell("Amount", (250, 50, 350, 70), 0, 1, is_header=True)
    h3 = TableCell("Rate", (350, 50, 450, 70), 0, 2, is_header=True)
    r0 = TableRow([h1, h2, h3], row_idx=0, is_header=True)

    c1 = TableCell("Cash Equivalents", (58, 75, 250, 95), 1, 0, indent_level=1)
    c2 = TableCell("$15,400.50", (250, 75, 350, 95), 1, 1, data_type="currency", typed_value=15400.50, format_code='"$"#,##0.00')
    c3 = TableCell("4.5%", (350, 75, 450, 95), 1, 2, data_type="percentage", typed_value=0.045, format_code="0.0%")
    r1 = TableRow([c1, c2, c3], row_idx=1)

    c4 = TableCell("Total Assets", (50, 100, 250, 120), 2, 0)
    c5 = TableCell("$15,400.50", (250, 100, 350, 120), 2, 1, data_type="currency", typed_value=15400.50, format_code='"$"#,##0.00')
    c6 = TableCell("100%", (350, 100, 450, 120), 2, 2, data_type="percentage", typed_value=1.0, format_code="0%")
    r2 = TableRow([c4, c5, c6], row_idx=2)

    table = TableBlock(bbox=(50, 50, 450, 120), rows=[r0, r1, r2])

    # 1. Test in-memory byte generation
    xlsx_bytes = writer.generate([table])
    assert isinstance(xlsx_bytes, bytes)
    assert len(xlsx_bytes) > 1000

    # Load and inspect openpyxl workbook
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb.active
    assert ws.title == "Extracted Table"

    # Verify Header cell styling
    cell_h1 = ws.cell(row=1, column=1)
    assert cell_h1.value == "Financial Line Item"
    assert cell_h1.font.bold is True

    # Verify Currency native float typing
    cell_amt = ws.cell(row=2, column=2)
    assert isinstance(cell_amt.value, float)
    assert cell_amt.value == 15400.50
    assert '"$"#,##0.00' in cell_amt.number_format

    # Verify Percentage typing
    cell_rate = ws.cell(row=2, column=3)
    assert isinstance(cell_rate.value, float)
    assert cell_rate.value == 0.045
    assert "%" in cell_rate.number_format

    # Verify Indentation
    cell_indented = ws.cell(row=2, column=1)
    assert cell_indented.alignment.indent == 1

    # Verify Auto Column Width
    assert ws.column_dimensions["A"].width >= 15

    # 2. Test file generation to disk
    out_file = str(tmp_path / "output_test.xlsx")
    saved_path = writer.write_to_path([table], out_file)
    assert os.path.exists(saved_path)
    assert os.path.getsize(saved_path) > 1000


def test_excel_writer_merged_cells():
    writer = EnterpriseExcelWriter()

    # Cell spanning 3 columns: colspan=3
    title_cell = TableCell("Consolidated Statement", (50, 50, 350, 80), 0, 0, colspan=3, is_header=True)
    r0 = TableRow([title_cell], row_idx=0)

    c1 = TableCell("A", (50, 85, 150, 105), 1, 0)
    c2 = TableCell("B", (150, 85, 250, 105), 1, 1)
    c3 = TableCell("C", (250, 85, 350, 105), 1, 2)
    r1 = TableRow([c1, c2, c3], row_idx=1)

    table = TableBlock(bbox=(50, 50, 350, 105), rows=[r0, r1], has_merged_cells=True)

    xlsx_bytes = writer.generate([table])
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb.active

    # Check that merge range exists
    merge_ranges = [str(r) for r in ws.merged_cells.ranges]
    assert any("A1:C1" in m for m in merge_ranges)


# =====================================================================
# 3. Container Registration Tests
# =====================================================================

def test_container_validator_and_writer_resolution():
    container = get_default_container()

    val = container.get_validator()
    assert isinstance(val, TableValidator)

    wr = container.get_excel_writer()
    assert isinstance(wr, EnterpriseExcelWriter)
