import os
import shutil
import pytest
import openpyxl
import fitz  # PyMuPDF

from app.services.engine.pdf_to_excel import (
    PdfToExcelConverter,
    detect_cell_type,
)
from app.services.engine.registry import converter_registry
from app.core.errors import FileValidationError

TEST_TMP_DIR = "./temp_pdf_excel_tests"


@pytest.fixture(autouse=True)
def setup_teardown():
    os.makedirs(TEST_TMP_DIR, exist_ok=True)
    yield
    if os.path.exists(TEST_TMP_DIR):
        shutil.rmtree(TEST_TMP_DIR, ignore_errors=True)


def create_sample_pdf_with_table(path: str) -> str:
    """Create a PDF with drawn table lines and text cells."""
    doc = fitz.open()
    page = doc.new_page(width=612, height=792)  # Standard Letter

    # Title
    page.insert_text((50, 50), "Quarterly Financial Statement", fontsize=16)

    # Draw table borders
    # 4 columns, 4 rows
    x_coords = [50, 180, 300, 420, 550]
    y_coords = [100, 130, 160, 190, 220]

    for y in y_coords:
        page.draw_line((x_coords[0], y), (x_coords[-1], y))
    for x in x_coords:
        page.draw_line((x, y_coords[0]), (x, y_coords[-1]))

    # Headers
    page.insert_text((55, 120), "Department", fontsize=11)
    page.insert_text((185, 120), "Q1 Revenue", fontsize=11)
    page.insert_text((305, 120), "Growth %", fontsize=11)
    page.insert_text((425, 120), "Status", fontsize=11)

    # Row 1
    page.insert_text((55, 150), "Enterprise Cloud", fontsize=10)
    page.insert_text((185, 150), "$1,450,200.50", fontsize=10)
    page.insert_text((305, 150), "18.5%", fontsize=10)
    page.insert_text((425, 150), "Active", fontsize=10)

    # Row 2
    page.insert_text((55, 180), "Consumer SaaS", fontsize=10)
    page.insert_text((185, 180), "$850,000.00", fontsize=10)
    page.insert_text((305, 180), "24.0%", fontsize=10)
    page.insert_text((425, 180), "Active", fontsize=10)

    # Row 3
    page.insert_text((55, 210), "Professional Svcs", fontsize=10)
    page.insert_text((185, 210), "$320,500.25", fontsize=10)
    page.insert_text((305, 210), "5.2%", fontsize=10)
    page.insert_text((425, 210), "Pending", fontsize=10)

    doc.save(path)
    doc.close()
    return path


def test_registry_registration():
    """Verify converter is registered in the global registry."""
    converter = converter_registry.get("pdf-to-excel")
    assert converter is not None
    assert converter.tool_id == "pdf-to-excel"
    assert "pdf" in converter.supported_inputs
    assert converter.output_extension == "xlsx"


def test_data_type_detection():
    """Test smart cell type inference for currency, percentages, integers, and floats."""
    val, num_format = detect_cell_type("$1,450,200.50")
    assert val == 1450200.50
    assert "$" in (num_format or "")

    val, num_format = detect_cell_type("18.5%")
    assert pytest.approx(val, 0.001) == 0.185
    assert "%" in (num_format or "")

    val, _ = detect_cell_type("42")
    assert val == 42

    val, _ = detect_cell_type("1,234.56")
    assert val == 1234.56

    val, _ = detect_cell_type("(500.00)")
    assert val == -500.00

    val, _ = detect_cell_type("Regular Text String")
    assert val == "Regular Text String"


@pytest.mark.asyncio
async def test_pdf_to_excel_conversion():
    """Verify full conversion pipeline from PDF to XLSX."""
    pdf_path = os.path.join(TEST_TMP_DIR, "sample_financial.pdf")
    create_sample_pdf_with_table(pdf_path)

    converter = PdfToExcelConverter()
    output_dir = os.path.join(TEST_TMP_DIR, "output")
    os.makedirs(output_dir, exist_ok=True)

    result = await converter.convert(
        input_paths=[pdf_path],
        output_dir=output_dir,
        options={"output_format": "xlsx", "ocr": False}
    )

    assert os.path.exists(result.output_path)
    assert result.output_path.endswith(".xlsx")
    assert result.size_bytes > 0

    # Verify workbook validity with openpyxl
    wb = openpyxl.load_workbook(result.output_path)
    assert len(wb.sheetnames) >= 1

    ws = wb.active
    assert ws is not None

    # Check that rows were populated
    rows = list(ws.iter_rows(values_only=True))
    assert len(rows) >= 1

    wb.close()


def test_api_tools_endpoint_includes_pdf_to_excel(client):
    """Verify that /api/v1/tools lists pdf-to-excel with correct metadata."""
    res = client.get("/api/v1/tools")
    assert res.status_code == 200
    data = res.json()
    assert data["success"] is True
    tool_ids = [t["tool_id"] for t in data["data"]]
    assert "pdf-to-excel" in tool_ids

    pdf_to_excel_tool = next(t for t in data["data"] if t["tool_id"] == "pdf-to-excel")
    assert pdf_to_excel_tool["name"] == "PDF to Excel"
    assert "pdf" in pdf_to_excel_tool["supported_inputs"]
    assert pdf_to_excel_tool["output_extension"] == "xlsx"

