"""
STEP 1 & STEP 2: Diagnostic Reproduction & Isolation Script for PDF-to-Word
Measures:
- page number
- scanned/native classification
- rasterization start/end
- image dimensions and DPI
- OCR start/end
- OCR duration
- RSS memory before rasterization
- RSS after rasterization
- RSS after OCR
- RSS after cleanup
- Word reconstruction start/end
- exception/exit if any
"""

import os
import sys
import time
import asyncio
import psutil
import fitz
from docx import Document
from docx.shared import Inches, Pt

backend_dir = os.path.dirname(os.path.abspath(__file__))
if backend_dir not in sys.path:
    sys.path.insert(0, backend_dir)

from app.core.ocr_guard import find_tesseract_bin, is_near_blank_page, ocr_resource_guard
import pytesseract
from PIL import Image

tess_bin = find_tesseract_bin()
if tess_bin:
    pytesseract.pytesseract.tesseract_cmd = tess_bin


def run_page_diagnostic(page_idx: int, page: fitz.Page, docx_out: Document, proc: psutil.Process):
    report = {
        "page_number": page_idx + 1,
        "scanned_classification": "unknown",
        "dimensions": None,
        "dpi": 200,
        "raster_duration": 0.0,
        "ocr_duration": 0.0,
        "reconstruct_duration": 0.0,
        "rss_before_raster": 0.0,
        "rss_after_raster": 0.0,
        "rss_after_ocr": 0.0,
        "rss_after_cleanup": 0.0,
        "exception": None,
        "skipped_blank": False,
    }

    try:
        raw_text = page.get_text().strip()
        is_scanned = len(raw_text) < 40
        report["scanned_classification"] = "scanned" if is_scanned else "native"

        if page_idx > 0:
            docx_out.add_page_break()

        if not is_scanned:
            extracted_text = raw_text
            t_rec_start = time.perf_counter()
            p = docx_out.add_paragraph()
            p.add_run(extracted_text)
            report["reconstruct_duration"] = time.perf_counter() - t_rec_start
            return report

        # Check blank page
        if is_near_blank_page(page):
            report["skipped_blank"] = True
            report["rss_before_raster"] = proc.memory_info().rss / (1024 * 1024)
            report["rss_after_cleanup"] = proc.memory_info().rss / (1024 * 1024)
            return report

        # Scanned page processing
        report["rss_before_raster"] = proc.memory_info().rss / (1024 * 1024)
        t_rast_start = time.perf_counter()
        pix = page.get_pixmap(dpi=200, alpha=False)
        report["raster_duration"] = time.perf_counter() - t_rast_start
        report["dimensions"] = (pix.width, pix.height)
        report["rss_after_raster"] = proc.memory_info().rss / (1024 * 1024)

        pil_img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
        del pix
        pix = None

        t_ocr_start = time.perf_counter()
        with ocr_resource_guard(job_id=f"diag_p{page_idx+1}", page_idx=page_idx, engine="tesseract_text"):
            extracted_text = pytesseract.image_to_string(pil_img, lang="eng")
        report["ocr_duration"] = time.perf_counter() - t_ocr_start
        report["rss_after_ocr"] = proc.memory_info().rss / (1024 * 1024)

        # Cleanup image
        pil_img.close()
        del pil_img
        pil_img = None
        import gc
        gc.collect()
        report["rss_after_cleanup"] = proc.memory_info().rss / (1024 * 1024)

        # Word reconstruction
        t_rec_start = time.perf_counter()
        blocks = extracted_text.split("\n\n")
        for block in blocks:
            clean_block = block.strip()
            if clean_block:
                p = docx_out.add_paragraph()
                lines = clean_block.splitlines()
                for line_idx, line in enumerate(lines):
                    run = p.add_run(line.strip())
                    run.font.name = "Calibri"
                    run.font.size = Pt(11)
                    if line_idx < len(lines) - 1:
                        p.add_run("\n")
        report["reconstruct_duration"] = time.perf_counter() - t_rec_start

    except Exception as e:
        report["exception"] = str(e)

    return report


def test_individual_pages(doc_path: str):
    print("\n" + "="*70)
    print("STEP 2A: TESTING PAGES INDIVIDUALLY (1 to 10)")
    print("="*70)
    doc = fitz.open(doc_path)
    proc = psutil.Process()

    results = []
    for p_idx in range(len(doc)):
        page = doc[p_idx]
        docx_out = Document()
        rep = run_page_diagnostic(p_idx, page, docx_out, proc)
        results.append(rep)
        print(f"Page {rep['page_number']}: class={rep['scanned_classification']}, blank={rep['skipped_blank']}, "
              f"dim={rep['dimensions']}, rast={rep['raster_duration']:.2f}s, ocr={rep['ocr_duration']:.2f}s, "
              f"rss_before={rep['rss_before_raster']:.1f}MB, rss_after_rast={rep['rss_after_raster']:.1f}MB, "
              f"rss_after_ocr={rep['rss_after_ocr']:.1f}MB, rss_after_clean={rep['rss_after_cleanup']:.1f}MB, "
              f"exc={rep['exception']}")
    doc.close()
    return results


def test_progressive_pages(doc_path: str):
    print("\n" + "="*70)
    print("STEP 2B: TESTING PAGES PROGRESSIVELY (1-2, 1-3, ..., 1-10)")
    print("="*70)
    proc = psutil.Process()

    for count in range(2, 11):
        doc = fitz.open(doc_path)
        docx_out = Document()
        rss_start = proc.memory_info().rss / (1024 * 1024)
        peak_rss = rss_start
        t0 = time.perf_counter()
        failed_page = None
        exception_str = None

        for p_idx in range(count):
            page = doc[p_idx]
            rep = run_page_diagnostic(p_idx, page, docx_out, proc)
            curr_rss = proc.memory_info().rss / (1024 * 1024)
            if curr_rss > peak_rss:
                peak_rss = curr_rss
            if rep["exception"]:
                failed_page = p_idx + 1
                exception_str = rep["exception"]
                break

        out_path = f"d:/convertlytools/backend/temp_prog_{count}.docx"
        if not failed_page:
            docx_out.save(out_path)
            size = os.path.getsize(out_path)
            os.remove(out_path)
        else:
            size = 0

        doc.close()
        t1 = time.perf_counter()
        rss_end = proc.memory_info().rss / (1024 * 1024)

        status = f"FAILED on page {failed_page} ({exception_str})" if failed_page else f"SUCCESS (size={size:,} bytes)"
        print(f"Pages 1-{count}: {status} | Duration={t1-t0:.2f}s | Peak RSS={peak_rss:.1f}MB | End RSS={rss_end:.1f}MB")


def main():
    doc_path = r"D:\convertlytools\Bank July & Aug 2026.pdf"
    test_individual_pages(doc_path)
    test_progressive_pages(doc_path)


if __name__ == "__main__":
    main()
